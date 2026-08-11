from __future__ import annotations

import json
import shutil
from copy import copy
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.chart.axis import DateAxis
from openpyxl.comments import Comment
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo


APP_ROOT = Path(r"F:\HBVStudio")
SOURCE_NAME = "正式_YC-ML_小时尺度_2025"
TARGET_NAME = "测试_YC-ML_小时尺度_洪水事件_2025"
SOURCE_RUNTIME = APP_ROOT / "用户数据" / "运行目录" / SOURCE_NAME
TARGET_RUNTIME = APP_ROOT / "用户数据" / "运行目录" / TARGET_NAME
SOURCE_CONFIG = APP_ROOT / "用户数据" / "workspaces" / f"{SOURCE_NAME}.json"
TARGET_CONFIG = APP_ROOT / "用户数据" / "workspaces" / f"{TARGET_NAME}.json"
SOURCE_OBS = SOURCE_RUNTIME / "数据" / "观测数据" / "ML出口流量.xlsx"
EVENT_BOOK = TARGET_RUNTIME / "数据" / "观测数据" / "洪水事件表_2025.xlsx"
PARAM_FILE = TARGET_RUNTIME / "数据" / "观测数据" / "连续率定最优参数_20260715_175134.json"

EVENTS = [
    {
        "event_id": "E20250601_CAL",
        "purpose": "calibration",
        "boundary_start": "2025-05-01 08:00",
        "score_start": "2025-05-29 19:00",
        "score_end": "2025-06-10 15:00",
        "note": "初夏涨水过程，率定场次。",
    },
    {
        "event_id": "E20250705_CAL",
        "purpose": "calibration",
        "boundary_start": "2025-05-01 08:00",
        "score_start": "2025-07-02 05:00",
        "score_end": "2025-07-13 12:00",
        "note": "7月洪水过程，率定场次。",
    },
    {
        "event_id": "E20250804_VAL",
        "purpose": "validation",
        "boundary_start": "2025-05-01 08:00",
        "score_start": "2025-07-31 10:00",
        "score_end": "2025-08-09 05:00",
        "note": "8月上旬洪水过程，验证场次。",
    },
    {
        "event_id": "E20250817_CAL",
        "purpose": "calibration",
        "boundary_start": "2025-05-01 08:00",
        "score_start": "2025-08-13 03:00",
        "score_end": "2025-08-21 13:00",
        "note": "8月中旬洪水过程，率定场次。",
    },
    {
        "event_id": "E20250902_CAL",
        "purpose": "calibration",
        "boundary_start": "2025-05-01 08:00",
        "score_start": "2025-08-28 11:00",
        "score_end": "2025-09-10 11:00",
        "note": "年内最大洪峰过程，率定场次。",
    },
    {
        "event_id": "E20250922_VAL",
        "purpose": "validation",
        "boundary_start": "2025-05-01 08:00",
        "score_start": "2025-09-17 17:00",
        "score_end": "2025-09-30 17:00",
        "note": "秋季高流量洪水过程，验证场次。",
    },
]


def source_boundary_path() -> Path:
    config = json.loads(SOURCE_CONFIG.read_text(encoding="utf-8"))
    raw = str(dict(config.get("边界条件", {}) or {}).get("上游边界入流_csv", "") or "").strip()
    if not raw:
        raise ValueError("源工作区未配置上游边界入流。")
    return Path(raw)


def formal_workspace_inventory() -> list[str]:
    return sorted(path.name for path in SOURCE_CONFIG.parent.glob("正式_*.json"))

BEST_PARAMS = {
    "TT": -1.701981,
    "FC": 103.213568,
    "BETA": 0.502183,
    "LP": 0.946249,
    "RFCF": 1.199056,
    "SFCF": 1.267366,
    "CFR": 0.063416,
    "CWH": 0.099486,
    "CFMAX_low": 4.748314,
    "CFMAX_high": 5.547941,
    "K": 0.485595,
    "K1": 0.125653,
    "K2": 0.016529,
    "UZL": 109.517935,
    "PERC": 0.019912,
    "ICE_FACTOR": 3.934206,
    "K_MUSK": 1.110856,
    "X_MUSK": 0.017239,
}


def portable(*parts: str) -> str:
    suffix = "/".join(parts)
    return f"__PROJECT_ROOT__/{suffix}"


def copy_tree_resume(source: Path, target: Path) -> None:
    target.mkdir(parents=True, exist_ok=True)
    for item in source.rglob("*"):
        relative = item.relative_to(source)
        destination = target / relative
        if item.is_dir():
            destination.mkdir(parents=True, exist_ok=True)
            continue
        destination.parent.mkdir(parents=True, exist_ok=True)
        if destination.exists() and destination.stat().st_size == item.stat().st_size:
            continue
        shutil.copy2(item, destination)


def copy_inputs() -> None:
    (TARGET_RUNTIME / "数据" / "观测数据").mkdir(parents=True, exist_ok=True)
    copy_tree_resume(SOURCE_RUNTIME / "数据" / "地理数据", TARGET_RUNTIME / "数据" / "地理数据")
    copy_tree_resume(SOURCE_RUNTIME / "数据" / "模型输入", TARGET_RUNTIME / "数据" / "模型输入")
    shutil.copy2(SOURCE_OBS, TARGET_RUNTIME / "数据" / "观测数据" / SOURCE_OBS.name)
    source_boundary = source_boundary_path()
    shutil.copy2(source_boundary, TARGET_RUNTIME / "数据" / "观测数据" / source_boundary.name)


def replace_workspace_name(value):
    if isinstance(value, dict):
        return {key: replace_workspace_name(item) for key, item in value.items()}
    if isinstance(value, list):
        return [replace_workspace_name(item) for item in value]
    if isinstance(value, str):
        return value.replace(SOURCE_NAME, TARGET_NAME)
    return value


def rewrite_copied_manifests() -> None:
    hourly = TARGET_RUNTIME / "数据" / "模型输入" / "小时尺度"
    for name in ("hourly_forcing_manifest.json", "hourly_forcing_summary.json", "_meteo_state.json"):
        path = hourly / name
        if not path.exists():
            continue
        data = json.loads(path.read_text(encoding="utf-8"))
        path.write_text(json.dumps(replace_workspace_name(data), ensure_ascii=False, indent=2), encoding="utf-8")


def style_sheet(ws, widths: dict[str, float]) -> None:
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False
    for key, width in widths.items():
        ws.column_dimensions[key].width = width
    for cell in ws[1]:
        cell.font = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 25
    thin = Side(style="thin", color="D9E2F3")
    for row in ws.iter_rows():
        for cell in row:
            font = copy(cell.font)
            font.name = "微软雅黑"
            font.sz = 10
            cell.font = font
            cell.alignment = Alignment(vertical="center")
            cell.border = Border(bottom=thin)


def add_table(ws, name: str) -> None:
    table = Table(displayName=name, ref=ws.dimensions)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
        showRowStripes=True, showColumnStripes=False,
    )
    ws.add_table(table)


def build_event_workbook() -> dict[str, object]:
    obs = pd.read_excel(SOURCE_OBS)
    obs.columns = ["date", "discharge_m3s"]
    obs["date"] = pd.to_datetime(obs["date"])

    wb = Workbook()
    ws = wb.active
    ws.title = "洪水事件表"
    headers = ["event_id", "purpose", "boundary_start", "score_start", "score_end", "note"]
    ws.append(headers)
    for event in EVENTS:
        ws.append([
            pd.Timestamp(event[key]).to_pydatetime() if key in {"boundary_start", "score_start", "score_end"} else event[key]
            for key in headers
        ])
    for row in ws.iter_rows(min_row=2, min_col=3, max_col=5):
        for cell in row:
            cell.number_format = "yyyy-mm-dd hh:mm"
    style_sheet(ws, {"A": 22, "B": 15, "C": 20, "D": 20, "E": 20, "F": 20})
    add_table(ws, "FloodEvents")
    purpose_validation = DataValidation(type="list", formula1='"calibration,validation,diagnostic"')
    ws.add_data_validation(purpose_validation)
    purpose_validation.add(f"B2:B{ws.max_row}")
    for cell in ws[1]:
        cell.comment = Comment("HBV-Studio 场次洪水接口字段，请勿改名。", "HBV-Studio")

    process = wb.create_sheet("场次实测过程")
    process.append(["event_id", "purpose", "date", "discharge_m3s"])
    stats: list[dict[str, object]] = []
    for event in EVENTS:
        mask = obs["date"].between(event["score_start"], event["score_end"])
        part = obs.loc[mask].copy()
        if part.empty or part["discharge_m3s"].isna().any():
            raise ValueError(f"事件评分窗口观测不完整：{event['event_id']}")
        peak_row = part.loc[part["discharge_m3s"].idxmax()]
        peak_position = int(part.index.get_loc(peak_row.name))
        if peak_position in {0, len(part) - 1}:
            raise ValueError(f"场次洪峰位于评价窗口边界：{event['event_id']}")
        for row in part.itertuples(index=False):
            process.append([event["event_id"], event["purpose"], row.date.to_pydatetime(), float(row.discharge_m3s)])
        stats.append(
            {
                "event_id": event["event_id"],
                "purpose": event["purpose"],
                "score_start": pd.Timestamp(event["score_start"]),
                "score_end": pd.Timestamp(event["score_end"]),
                "hours": int(len(part)),
                "peak_time": pd.Timestamp(peak_row["date"]),
                "peak_flow_m3s": float(peak_row["discharge_m3s"]),
                "start_flow_m3s": float(part.iloc[0]["discharge_m3s"]),
                "end_flow_m3s": float(part.iloc[-1]["discharge_m3s"]),
                "observed_volume_1e8m3": float(part["discharge_m3s"].sum() * 3600.0 / 1e8),
                "coverage": "100%",
            }
        )
    for row in process.iter_rows(min_row=2, min_col=3, max_col=3):
        row[0].number_format = "yyyy-mm-dd hh:mm"
    for row in process.iter_rows(min_row=2, min_col=4, max_col=4):
        row[0].number_format = "0.0"
    style_sheet(process, {"A": 22, "B": 15, "C": 20, "D": 20})
    add_table(process, "ObservedEventHydrographs")
    process.conditional_formatting.add(
        f"D2:D{process.max_row}",
        ColorScaleRule(start_type="min", start_color="E2F0D9", mid_type="percentile", mid_value=50,
                       mid_color="FFF2CC", end_type="max", end_color="F4CCCC"),
    )

    summary = wb.create_sheet("事件统计")
    summary_headers = [
        "event_id", "purpose", "score_start", "score_end", "hours", "peak_time",
        "peak_flow_m3s", "start_flow_m3s", "end_flow_m3s", "observed_volume_1e8m3", "coverage",
    ]
    summary.append(summary_headers)
    for item in stats:
        summary.append([item[key].to_pydatetime() if isinstance(item[key], pd.Timestamp) else item[key] for key in summary_headers])
    for row in summary.iter_rows(min_row=2, min_col=3, max_col=6):
        for cell in row:
            if isinstance(cell.value, pd.Timestamp):
                cell.value = cell.value.to_pydatetime()
            if hasattr(cell.value, "year"):
                cell.number_format = "yyyy-mm-dd hh:mm"
    for row in summary.iter_rows(min_row=2, min_col=7, max_col=10):
        for cell in row:
            cell.number_format = "0.000"
    style_sheet(summary, {"A": 22, "B": 15, "C": 20, "D": 20, "E": 10, "F": 20,
                          "G": 18, "H": 18, "I": 18, "J": 24, "K": 12})
    add_table(summary, "EventStatistics")

    chart = LineChart()
    chart.title = "ML 出口洪水场次实测过程"
    chart.style = 13
    chart.y_axis.title = "流量 (m3/s)"
    chart.x_axis = DateAxis(crosses="autoZero")
    chart.x_axis.title = "时间"
    chart.x_axis.number_format = "yyyy-mm-dd"
    chart.height = 9
    chart.width = 24
    chart.add_data(Reference(process, min_col=4, min_row=1, max_row=process.max_row), titles_from_data=True)
    chart.set_categories(Reference(process, min_col=3, min_row=2, max_row=process.max_row))
    summary.add_chart(chart, "A6")

    notes = wb.create_sheet("使用说明")
    notes.sheet_view.showGridLines = False
    note_rows = [
        ("用途", "本文件首张工作表可由 HBV-Studio 直接读取；其余工作表用于人工复核。"),
        ("场次划分", "以实测流量过程识别候选洪峰，结合降水和气温过程复核，并以涨水前及主要退水后的局部低点确定评价窗口。"),
        ("率定场次", "E20250601_CAL、E20250705_CAL、E20250817_CAL 和 E20250902_CAL 进入 flood_event_calibration_v1 目标函数，各场次等权。"),
        ("验证场次", "E20250804_VAL 和 E20250922_VAL 不参与参数寻优，仅使用最优参数独立评价。"),
        ("模型预热期", "模型自 2025-01-01 08:00 连续演算，积雪水当量、土壤含水量及上下层响应库状态在场次之间连续传递。"),
        ("边界汇流预热期", "边界入流自 2025-05-01 08:00 连续有测；缺测时段不按零流量处理，资料恢复后设置14天边界汇流预热期。"),
        ("适用限制", "案例仅覆盖2025年6场洪水，能够演示独立验证流程，但不代表已覆盖所有洪峰等级或产流成因。"),
    ]
    notes.append(["项目", "说明"])
    for row in note_rows:
        notes.append(list(row))
    style_sheet(notes, {"A": 18, "B": 105})
    for row in notes.iter_rows(min_row=2, min_col=2, max_col=2):
        row[0].alignment = Alignment(vertical="top", wrap_text=True)
    notes.row_dimensions[2].height = 34
    notes.row_dimensions[3].height = 46
    add_table(notes, "WorkbookNotes")

    wb.save(EVENT_BOOK)
    return {"events": stats, "process_rows": process.max_row - 1}


def build_config() -> None:
    config = json.loads(SOURCE_CONFIG.read_text(encoding="utf-8"))
    data_root = TARGET_RUNTIME / "数据"
    config.update(
        {
            "目标函数模式": "flood_event_calibration_v1",
            "任务时段模式": "continuous",
            "场次洪水工作流": "continuous_events",
            "运行目录": str(TARGET_RUNTIME),
            "流域名称": TARGET_NAME,
            "流域编号": TARGET_NAME,
            "流域边界_shp": str(data_root / "地理数据" / "basin.shp"),
            "DEM_tif": str(data_root / "地理数据" / "dem_0p1deg.tif"),
            "冰川边界_shp": str(data_root / "地理数据" / "glacier_shp" / "glacier.shp"),
            "观测径流_csv": str(data_root / "观测数据" / "ML出口流量.xlsx"),
            "时间": {
                "开始年份": 2025,
                "结束年份": 2025,
                "预热开始": "2025-01-01 08:00",
                "预热结束": "2025-05-15 07:00",
                "率定开始": "2025-05-15 08:00",
                "率定结束": "2025-09-10 11:00",
                "验证开始": "2025-09-17 17:00",
                "验证结束": "2025-09-30 17:00",
            },
            "事件资料模式": {
                "启用": False,
                "事件窗口资料": False,
                "事件表路径": str(EVENT_BOOK),
                "允许事件间断": False,
                "初始条件策略": "continuous_state",
            },
            "洪水事件率定": {
                "启用": True,
                "事件窗口资料": False,
                "事件表路径": str(EVENT_BOOK),
                "模式": "objective",
                "峰现容许误差小时": 6,
                "边界汇流预热天数": 14,
                "初始条件策略": "continuous_state",
                "目标事件类型": ["calibration"],
                "目标权重": {
                    "洪峰流量误差": 0.30,
                    "峰现时间误差": 0.25,
                    "洪量误差": 0.25,
                    "退水过程误差": 0.10,
                    "高流量过程效率": 0.10,
                },
            },
            "边界条件": {
                "上游边界入流_csv": str(data_root / "观测数据" / "YC上游边界入流.csv"),
                "时间字段": "date",
                "流量字段": "flow",
                "缺失填补": "preserve_missing",
            },
        }
    )
    meteo = dict(config.get("气象策略", {}))
    meteo["自带降水tif目录"] = str(data_root / "模型输入" / "小时尺度" / "降水_本地导入")
    meteo["自带温度tif目录"] = str(data_root / "模型输入" / "小时尺度" / "气温")
    meteo["自带蒸散发tif目录"] = str(data_root / "模型输入" / "小时尺度" / "蒸散发")
    config["气象策略"] = meteo
    TARGET_CONFIG.write_text(json.dumps(config, ensure_ascii=False, indent=2), encoding="utf-8")


def verify_outputs(expected_process_rows: int) -> None:
    inventory = formal_workspace_inventory()
    if len(inventory) != 6:
        raise ValueError(f"预期识别6个“正式_”工作区，实际为 {len(inventory)}：{inventory}")
    wb = load_workbook(EVENT_BOOK, data_only=False)
    if wb.sheetnames[:4] != ["洪水事件表", "场次实测过程", "事件统计", "使用说明"]:
        raise ValueError(f"工作表顺序不正确：{wb.sheetnames}")
    headers = [cell.value for cell in wb["洪水事件表"][1]]
    if headers != ["event_id", "purpose", "boundary_start", "score_start", "score_end", "note"]:
        raise ValueError(f"事件表字段不正确：{headers}")
    if wb["场次实测过程"].max_row - 1 != expected_process_rows:
        raise ValueError("场次实测过程行数校验失败。")
    if len(wb["事件统计"]._charts) != 1:
        raise ValueError("事件统计图表未生成。")
    formula_errors = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and any(err in cell.value for err in ("#REF!", "#DIV/0!", "#VALUE!", "#N/A", "#NAME?")):
                    formula_errors.append(f"{ws.title}!{cell.coordinate}={cell.value}")
    if formula_errors:
        raise ValueError("Excel 错误：" + "; ".join(formula_errors[:5]))
    config = json.loads(TARGET_CONFIG.read_text(encoding="utf-8"))
    if config.get("任务时段模式") != "continuous" or config.get("目标函数模式") != "flood_event_calibration_v1":
        raise ValueError("工作区事件模式配置校验失败。")
    if config.get("场次洪水工作流") != "continuous_events":
        raise ValueError("工作区未采用连续状态多场洪水工作流。")
    if dict(config.get("边界条件", {})).get("缺失填补") != "preserve_missing":
        raise ValueError("边界缺测处理未设置为保留缺测。")
    event_frame = pd.read_excel(EVENT_BOOK, sheet_name="洪水事件表")
    if event_frame["event_id"].duplicated().any():
        raise ValueError("场次洪水编号存在重复。")
    ordered = event_frame.sort_values("score_start")
    for left, right in zip(ordered.to_dict("records"), ordered.iloc[1:].to_dict("records")):
        if pd.Timestamp(left["score_end"]) >= pd.Timestamp(right["score_start"]):
            raise ValueError(f"场次洪水评价窗口重叠：{left['event_id']} 与 {right['event_id']}")
    purpose_counts = event_frame["purpose"].value_counts().to_dict()
    if purpose_counts.get("calibration") != 4 or purpose_counts.get("validation") != 2:
        raise ValueError(f"率定/验证场次配置不正确：{purpose_counts}")


def main() -> None:
    copy_inputs()
    rewrite_copied_manifests()
    result = build_event_workbook()
    PARAM_FILE.write_text(
        json.dumps(
            {
                "_来源": "正式_YC-ML_小时尺度_2025 / hbv_cryo_custom_tif_inline_20260715_175134",
                "_用途": "洪水事件示例快速试跑的初始参数；不是本事件目标函数的正式率定成果。",
                **BEST_PARAMS,
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )
    build_config()
    verify_outputs(int(result["process_rows"]))
    print(json.dumps(
        {
            "workspace_config": str(TARGET_CONFIG),
            "runtime": str(TARGET_RUNTIME),
            "event_workbook": str(EVENT_BOOK),
            "initial_params": str(PARAM_FILE),
            "events": result["events"],
            "formal_workspaces": formal_workspace_inventory(),
        },
        ensure_ascii=False,
        indent=2,
        default=str,
    ))


if __name__ == "__main__":
    main()
